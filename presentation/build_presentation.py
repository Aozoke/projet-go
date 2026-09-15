from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation"
PPTX_PATH = OUT / "wasmredis-presentation_2.pptx"
PDF_PATH = OUT / "wasmredis-presentation(2).pdf"
NOTES_PATH = OUT / "notes-orales(2).md"

PPT_W = Inches(13.333)
PPT_H = Inches(7.5)
PDF_W = 960
PDF_H = 540

BG = (10, 16, 25)
PANEL = (18, 28, 41)
PANEL_LIGHT = (28, 41, 57)
TEXT = (239, 244, 247)
MUTED = (153, 169, 183)
TEAL = (45, 193, 183)
ORANGE = (245, 127, 77)
YELLOW = (244, 205, 96)
GREEN = (94, 205, 145)
RED = (240, 93, 109)


@dataclass
class Snippet:
    path: str
    start: int
    lines: list[str]
    line_numbers: list[int] | None = None

    @property
    def label(self) -> str:
        end = self.line_numbers[-1] if self.line_numbers else self.start + len(self.lines) - 1
        return f"{self.path}:{self.start}-{end}"

    @property
    def numbered(self) -> str:
        numbers = self.line_numbers or list(range(self.start, self.start + len(self.lines)))
        width = len(str(numbers[-1]))
        return "\n".join(
            f"{number:>{width}}  {line.expandtabs(4)}"
            for number, line in zip(numbers, self.lines)
        )


def snippet(path: str, marker: str, count: int) -> Snippet:
    lines = (ROOT / path).read_text(encoding="utf-8").splitlines()
    start = next(index for index, line in enumerate(lines) if marker in line)
    selected = lines[start : start + count]
    if len(selected) > 12:
        raise ValueError(f"Extrait trop long : {path}")
    return Snippet(path, start + 1, selected)


def compact_function(path: str, marker: str) -> Snippet:
    source = (ROOT / path).read_text(encoding="utf-8").splitlines()
    start = next(index for index, line in enumerate(source) if marker in line)
    selected: list[str] = []
    numbers: list[int] = []
    depth = 0

    for index, line in enumerate(source[start:], start):
        depth += line.count("{") - line.count("}")
        if line.strip() and not line.lstrip().startswith("//"):
            selected.append(line)
            numbers.append(index + 1)
        if depth == 0:
            break

    if len(selected) > 12:
        raise ValueError(f"Extrait trop long : {path}")
    return Snippet(path, start + 1, selected, numbers)


def picked_snippet(path: str, line_numbers: list[int]) -> Snippet:
    source = (ROOT / path).read_text(encoding="utf-8").splitlines()
    selected = [source[number - 1] for number in line_numbers]
    if len(selected) > 12:
        raise ValueError(f"Extrait trop long : {path}")
    return Snippet(path, line_numbers[0], selected, line_numbers)


SLIDES = [
    {
        "kind": "cover",
        "title": "WasmRedis",
        "subtitle": "Une base Redis lite, entierement dans le navigateur",
        "note": (
            "Je presente une base cle-valeur volontairement simple. Le coeur metier est en Go, "
            "compile en WebAssembly. React pilote le moteur sans serveur, et OPFS garde les donnees."
        ),
    },
    {
        "kind": "flow",
        "title": "Une commande, cinq etapes",
        "subtitle": "Le fil rouge de toute la presentation",
        "takeaways": [
            "React ne bloque jamais : le moteur et le disque vivent dans le worker.",
            "La RAM donne la vitesse ; AOF + snapshot donnent la persistance.",
        ],
        "note": (
            "Je pars d'un clic dans React. Le SDK valide puis envoie un message au worker. "
            "Le worker appelle Go/WASM. Go modifie le state et son buffer, puis le worker vide "
            "ce buffer dans OPFS. C'est la carte mentale a garder pour la suite."
        ),
    },
    {
        "kind": "bridge",
        "title": "React atteint le moteur Go",
        "subtitle": "SDK, Worker et WASM transportent la commande",
        "snippets": [compact_function("cmd/wasm/main.go", "func wasmRedisExecute")],
        "takeaways": [
            "Le SDK valide, puis le worker appelle le bridge Go/WASM.",
            "OPFS impose le Worker : ses Sync Access Handles n'existent pas sur le main thread.",
        ],
        "note": (
            "Le schema montre le trajet depuis React jusqu'au moteur. La fonction affichee est le "
            "point d'entree Go expose a JavaScript : wasmRedisExecute verifie la requete, decode le "
            "JSON puis appelle Engine.ExecuteBatch. Le moteur et les acces OPFS restent dans le "
            "worker car les Sync Access Handles ne sont pas disponibles sur le main thread."
        ),
    },
    {
        "kind": "code",
        "title": "Le parser protege l'entree",
        "subtitle": "Du texte vers une commande Go valide",
        "snippets": [picked_snippet("internal/redis/parser.go", [35, 38, 41, 44, 47, 48, 49, 50, 51, 66, 67, 68])],
        "takeaways": [
            "GET WHERE est distingue d'un GET par cle.",
            "Toute commande inconnue est rejetee avant le moteur.",
        ],
        "note": (
            "ParseCommand nettoie le texte puis lit le nom de commande. L'extrait garde la decision "
            "la plus interessante : pour GET, le parser distingue GET WHERE d'une lecture par cle. "
            "Le default montre aussi qu'une commande inconnue est rejetee avant le moteur."
        ),
    },
    {
        "kind": "engine",
        "title": "Du texte au state",
        "subtitle": "Le moteur choisit l'action et produit les ecritures",
        "snippets": [picked_snippet("internal/redis/engine.go", [181, 182, 183, 184, 192, 193, 201, 208, 209, 210, 217])],
        "takeaways": [
            "Le switch relie chaque CommandType a Set, Get ou Delete sur state.",
            "Les ecritures renvoient une Operation qui sera ajoutee au buffer.",
        ],
        "note": (
            "La fonction execute contient la decision importante : son switch choisit Set, Get ou "
            "Delete. Set et Delete modifient state et renvoient une Operation a persister, alors que "
            "Get renvoie seulement la valeur. Execute ajoutera ensuite ces operations au buffer. "
            "Les numeros de ligne sautent car seuls les trois cas utiles sont affiches."
        ),
    },
    {
        "kind": "persistence",
        "title": "AOF + snapshot : restaurer l'etat",
        "subtitle": "Le snapshot est charge avant le rejeu du journal",
        "snippets": [
            picked_snippet(
                "web/src/worker/wasmRedis.worker.ts",
                [310, 313, 314, 315, 316, 320, 321, 322, 324, 325, 331, 332],
            )
        ],
        "takeaways": [
            "Au demarrage : snapshot d'abord, puis operations AOF rejouees dans l'ordre.",
            "Tout vient de .env, ex. VITE_WASMREDIS_FLUSH_INTERVAL_MS=1000.",
        ],
        "note": (
            "restoreFromOpfs est la fonction centrale du redemarrage. Elle lit snapshot.json et le "
            "charge dans Go, puis lit aof.log, reconstruit les operations valides et les rejoue. "
            "Les lignes de validation secondaire sont coupees, mais chaque ligne visible vient du "
            "fichier indique. Les intervalles et tailles sont fournis par le fichier .env."
        ),
    },
    {
        "kind": "code",
        "title": "Un batch, un seul passage",
        "subtitle": "Les commandes restent ordonnees et sans interleaving",
        "snippets": [compact_function("internal/redis/engine.go", "func (engine *Engine) ExecuteBatch")],
        "takeaways": [
            "batchMu verrouille tout le lot et les resultats restent alignes aux commandes.",
            "Mesure du projet : 15,5x plus rapide pour 30 SET regroupes.",
        ],
        "note": (
            "ExecuteBatch prend un verrou pour qu'aucun autre batch ne s'intercale. La boucle execute "
            "chaque commande dans l'ordre, ajoute son resultat au meme index logique et rassemble les "
            "ecritures. Le benchmark mesure 15,5 fois moins de cout pour trente SET regroupes."
        ),
    },
    {
        "kind": "code",
        "title": "Les plages passent par le B-Tree",
        "subtitle": "Le parcours s'arrete des que la limite est depassee",
        "snippets": [snippet("internal/redis/btree.go", "func (tree *BTree) RangeItems", 12)],
        "takeaways": [
            "equals utilise un index inverse ; contains assume un scan.",
            "Une plage selective reste a environ 0,3 us jusqu'a 10 000 entrees.",
        ],
        "note": (
            "Pour une recherche inferieure, le B-Tree est parcouru dans l'ordre croissant. "
            "Des qu'une valeur ne correspond plus, la fonction retourne false et coupe le parcours. "
            "Le cas superieur fait la meme chose en sens inverse."
        ),
    },
    {
        "kind": "code",
        "title": "Le TTL supprime vraiment la cle",
        "subtitle": "State, index et AOF restent coherents",
        "snippets": [snippet("internal/redis/engine.go", "func (engine *Engine) removeIfExpiredLocked", 9)],
        "takeaways": [
            "Expiration lazy au GET, plus balayage actif periodique.",
            "Le restore garde ExpiresAt et ne remplit pas de nouveau le buffer.",
        ],
        "note": (
            "Une valeur garde une date ExpiresAt. Si elle est depassee, cette fonction retire la cle "
            "du state et des deux index. L'appelant cree ensuite une operation DELETE pour l'AOF. "
            "Une horloge injectable rend ce comportement testable sans attendre."
        ),
    },
    {
        "kind": "code",
        "title": "100 000 entrees, 29 lignes DOM",
        "subtitle": "Virtual scroll maison + abonnement cible",
        "snippets": [
            snippet("web/src/virtual/getVirtualRange.ts", "const firstVisible", 6),
            snippet("web/src/store/entryStore.ts", "export function useEntry(key", 5),
        ],
        "takeaways": [
            "Le spacer simule la hauteur totale ; seule la tranche visible est montee.",
            "Modifier une ligne : compteur 1 -> 2 ; autre ligne : 1 -> 1.",
        ],
        "note": (
            "Le calcul transforme scrollTop en startIndex et endIndex, avec huit lignes de marge. "
            "Le test utilise 100 000 entrees et ne depasse jamais 29 lignes. Chaque EntryRow s'abonne "
            "a une seule cle avec useSyncExternalStore, donc une edition ne reveille pas les autres."
        ),
    },
    {
        "kind": "results",
        "title": "Ce que la demo prouve",
        "subtitle": "Mesures reelles, parcours reproductible",
        "note": (
            "Je termine par une demo courte : ajouter une cle avec TTL, forcer le flush, recharger, "
            "lancer Seed 100, filtrer value >= 50, editer une ligne puis lancer Benchmark. "
            "Le dataset 100k existe dans le repo mais n'est pas charge par defaut sur ce petit PC."
        ),
    },
]


def ppt_color(value):
    return RGBColor(*value)


def pdf_color(value):
    return colors.Color(value[0] / 255, value[1] / 255, value[2] / 255)


def ppt_rect(slide, left, top, width, height, fill, radius=False, line=None):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = ppt_color(fill)
    if line:
        shape.line.color.rgb = ppt_color(line)
    else:
        shape.line.fill.background()
    return shape


def ppt_text(slide, text, left, top, width, height, size, color=TEXT, bold=False, font="Aptos"):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = 0
    frame.margin_bottom = 0
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = ppt_color(color)
    return box


def ppt_background(slide):
    ppt_rect(slide, 0, 0, PPT_W, PPT_H, BG)


def ppt_header(slide, data, index):
    ppt_text(slide, data["title"], Inches(0.55), Inches(0.28), Inches(10.8), Inches(0.42), 25, TEXT, True)
    ppt_text(slide, data["subtitle"], Inches(0.57), Inches(0.77), Inches(10.8), Inches(0.28), 11.5, MUTED)
    ppt_text(slide, f"{index:02d}", Inches(12.35), Inches(0.35), Inches(0.4), Inches(0.25), 10, MUTED, True)


def code_text(snippets):
    return "\n\n".join(item.numbered for item in snippets)


def code_size(snippets, width_inches, height_inches):
    lines = code_text(snippets).splitlines()
    longest = max(len(line) for line in lines)
    by_width = width_inches * 72 / max(longest * 0.60, 1)
    by_height = height_inches * 72 / max(len(lines) * 1.25, 1)
    return max(8.5, min(15, by_width, by_height))


def ppt_code_height(snippets):
    line_count = len(code_text(snippets).splitlines())
    return Inches(min(4.25, max(2.25, 0.92 + line_count * 0.22)))


def ppt_code_panel(slide, snippets, left, top, width, height):
    ppt_rect(slide, left, top, width, height, PANEL, True, PANEL_LIGHT)
    labels = "   |   ".join(item.label for item in snippets)
    ppt_text(slide, labels, left + Inches(0.28), top + Inches(0.20), width - Inches(0.56), Inches(0.24), 9.5, TEAL, True)
    ppt_text(
        slide,
        code_text(snippets),
        left + Inches(0.30),
        top + Inches(0.62),
        width - Inches(0.60),
        height - Inches(0.82),
        code_size(snippets, width.inches - 0.60, height.inches - 0.82),
        TEXT,
        False,
        "Courier New",
    )


def ppt_code(slide, snippets):
    height = ppt_code_height(snippets)
    top = Inches(1.40 + (4.68 - height.inches) / 2)
    ppt_code_panel(
        slide,
        snippets,
        Inches(0.55),
        top,
        Inches(12.22),
        height,
    )


def ppt_code_after_flow(slide, snippets):
    height = ppt_code_height(snippets)
    top = Inches(2.28 + (3.78 - height.inches) / 2)
    ppt_code_panel(slide, snippets, Inches(0.55), top, Inches(12.22), height)


def ppt_stage_flow(slide, labels, colors, top):
    left = Inches(0.55)
    total_width = 12.22
    gap = 0.28
    node_width = (total_width - gap * (len(labels) - 1)) / len(labels)

    for position, (label, color) in enumerate(zip(labels, colors)):
        node_left = left + Inches(position * (node_width + gap))
        ppt_rect(slide, node_left, top, Inches(node_width), Inches(0.70), PANEL_LIGHT, True)
        ppt_rect(slide, node_left, top, Inches(0.07), Inches(0.70), color)
        ppt_text(
            slide,
            label,
            node_left + Inches(0.14),
            top + Inches(0.22),
            Inches(node_width - 0.22),
            Inches(0.25),
            10.2,
            TEXT,
            True,
        )
        if position < len(labels) - 1:
            arrow_left = node_left + Inches(node_width + 0.07)
            ppt_text(slide, ">", arrow_left, top + Inches(0.20), Inches(0.14), Inches(0.24), 14, ORANGE, True)


def ppt_engine(slide, data, index):
    ppt_background(slide)
    ppt_header(slide, data, index)
    ppt_stage_flow(
        slide,
        ["Command", "execute(command)", "Set / Get / Delete", "state", "buffer"],
        [TEAL, ORANGE, GREEN, TEAL, RED],
        Inches(1.40),
    )
    ppt_code_after_flow(slide, data["snippets"])
    ppt_takeaways(slide, data["takeaways"])


def ppt_persistence(slide, data, index):
    ppt_background(slide)
    ppt_header(slide, data, index)
    ppt_stage_flow(
        slide,
        ["state", "buffer Go", "AOF", "snapshot", "restore"],
        [TEAL, YELLOW, ORANGE, RED, GREEN],
        Inches(1.40),
    )
    ppt_code_after_flow(slide, data["snippets"])
    ppt_takeaways(slide, data["takeaways"])


def ppt_bridge(slide, data, index):
    ppt_background(slide)
    ppt_header(slide, data, index)
    ppt_stage_flow(
        slide,
        ["React", "SDK", "postMessage", "Worker", "wasmRedisExecute", "Engine"],
        [TEAL, YELLOW, ORANGE, ORANGE, GREEN, TEAL],
        Inches(1.40),
    )
    ppt_code_after_flow(slide, data["snippets"])
    ppt_takeaways(slide, data["takeaways"])


def ppt_takeaways(slide, points):
    width = Inches(5.95)
    for index, point in enumerate(points[:2]):
        left = Inches(0.58) + index * Inches(6.18)
        ppt_rect(slide, left, Inches(6.42), Inches(0.08), Inches(0.48), TEAL if index == 0 else ORANGE)
        ppt_text(slide, point, left + Inches(0.20), Inches(6.40), width, Inches(0.48), 11.5, TEXT, True)


def ppt_cover(slide, data):
    ppt_background(slide)
    ppt_rect(slide, 0, 0, Inches(0.16), PPT_H, ORANGE)
    ppt_text(slide, "WASM / REDIS / REACT", Inches(0.78), Inches(0.70), Inches(5.6), Inches(0.25), 11, TEAL, True)
    ppt_text(slide, data["title"], Inches(0.76), Inches(1.18), Inches(7.0), Inches(0.82), 51, TEXT, True)
    ppt_text(slide, data["subtitle"], Inches(0.80), Inches(2.12), Inches(7.8), Inches(0.45), 18, MUTED)
    ppt_text(slide, "Go en RAM. OPFS sur disque. React reste fluide.", Inches(0.80), Inches(3.03), Inches(7.4), Inches(0.42), 19, TEXT, True)

    stages = [("React", TEAL), ("SDK TS", YELLOW), ("Worker", ORANGE), ("Go/WASM", GREEN), ("OPFS", RED)]
    for index, (name, color) in enumerate(stages):
        top = Inches(1.00 + index * 1.05)
        ppt_rect(slide, Inches(9.45), top, Inches(2.65), Inches(0.64), color, True)
        ppt_text(slide, name, Inches(9.72), top + Inches(0.16), Inches(2.1), Inches(0.25), 15, BG, True)
    ppt_text(slide, "Projet local, sans serveur", Inches(0.80), Inches(6.82), Inches(4.0), Inches(0.25), 10, MUTED)


def ppt_flow(slide, data, index):
    ppt_background(slide)
    ppt_header(slide, data, index)
    stages = [
        ("React", "CRUD + liste", TEAL),
        ("SDK", "types + Zod", YELLOW),
        ("Worker", "postMessage", ORANGE),
        ("Go/WASM", "state + index", GREEN),
        ("OPFS", "AOF + snapshot", RED),
    ]
    for pos, (name, detail, color) in enumerate(stages):
        left = Inches(0.55 + pos * 2.53)
        ppt_rect(slide, left, Inches(1.62), Inches(2.05), Inches(1.32), PANEL, True, PANEL_LIGHT)
        ppt_rect(slide, left, Inches(1.62), Inches(2.05), Inches(0.10), color)
        ppt_text(slide, name, left + Inches(0.18), Inches(2.00), Inches(1.7), Inches(0.28), 16, TEXT, True)
        ppt_text(slide, detail, left + Inches(0.18), Inches(2.40), Inches(1.7), Inches(0.22), 9.5, MUTED)
        if pos < 4:
            ppt_text(slide, ">", left + Inches(2.17), Inches(2.05), Inches(0.25), Inches(0.3), 18, ORANGE, True)

    ppt_text(slide, "Cycle de persistance", Inches(0.62), Inches(3.55), Inches(3.0), Inches(0.28), 13, MUTED, True)
    cycle = [("state RAM", TEAL), ("buffer Go", YELLOW), ("AOF ~1 s", ORANGE), ("snapshot ~2 min", RED)]
    for pos, (name, color) in enumerate(cycle):
        left = Inches(0.62 + pos * 3.10)
        ppt_rect(slide, left, Inches(4.05), Inches(2.52), Inches(0.78), PANEL_LIGHT, True)
        ppt_rect(slide, left, Inches(4.05), Inches(0.08), Inches(0.78), color)
        ppt_text(slide, name, left + Inches(0.25), Inches(4.29), Inches(2.0), Inches(0.25), 13, TEXT, True)
        if pos < 3:
            ppt_text(slide, ">", left + Inches(2.68), Inches(4.25), Inches(0.25), Inches(0.25), 18, color, True)
    ppt_takeaways(slide, data["takeaways"])


def ppt_results(slide, data, index):
    ppt_background(slide)
    ppt_header(slide, data, index)
    metrics = [
        ("15,5 x", "gain batch / 30 SET", TEAL),
        ("~0,3 us", "range selective / 10k", GREEN),
        ("60 FPS", "scroll mesure", YELLOW),
        ("29", "lignes DOM / 100k", ORANGE),
    ]
    for pos, (value, label, color) in enumerate(metrics):
        left = Inches(0.56 + pos * 3.13)
        ppt_rect(slide, left, Inches(1.55), Inches(2.77), Inches(1.42), PANEL, True, PANEL_LIGHT)
        ppt_text(slide, value, left + Inches(0.22), Inches(1.86), Inches(2.2), Inches(0.44), 25, color, True)
        ppt_text(slide, label, left + Inches(0.22), Inches(2.42), Inches(2.2), Inches(0.24), 10, MUTED)

    ppt_text(slide, "Demo en 60 secondes", Inches(0.62), Inches(3.55), Inches(4.0), Inches(0.32), 17, TEXT, True)
    demo = [
        "1  SET avec TTL, puis Flush",
        "2  Recharger : la cle revient depuis OPFS",
        "3  Seed 100, puis value >= 50",
        "4  Editer une ligne : seule elle passe de 1 a 2",
        "5  Benchmark : latences, restore et FPS",
    ]
    for pos, line in enumerate(demo):
        ppt_text(slide, line, Inches(0.68), Inches(4.10 + pos * 0.47), Inches(7.2), Inches(0.28), 13, TEXT, pos == 0)

    ppt_rect(slide, Inches(8.55), Inches(3.62), Inches(4.15), Inches(2.55), PANEL_LIGHT, True)
    ppt_text(slide, "Validation finale", Inches(8.85), Inches(3.93), Inches(3.4), Inches(0.3), 16, YELLOW, True)
    ppt_text(slide, "Go race tests", Inches(8.88), Inches(4.50), Inches(2.4), Inches(0.25), 13, TEXT, True)
    ppt_text(slide, "Tests TypeScript", Inches(8.88), Inches(4.95), Inches(2.4), Inches(0.25), 13, TEXT, True)
    ppt_text(slide, "Build WASM + React", Inches(8.88), Inches(5.40), Inches(2.8), Inches(0.25), 13, TEXT, True)
    for pos in range(3):
        ppt_text(slide, "OK", Inches(11.55), Inches(4.50 + pos * 0.45), Inches(0.55), Inches(0.25), 13, GREEN, True)


def build_pptx():
    presentation = Presentation()
    presentation.slide_width = PPT_W
    presentation.slide_height = PPT_H
    for index, data in enumerate(SLIDES, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        if data["kind"] == "cover":
            ppt_cover(slide, data)
        elif data["kind"] == "flow":
            ppt_flow(slide, data, index)
        elif data["kind"] == "engine":
            ppt_engine(slide, data, index)
        elif data["kind"] == "persistence":
            ppt_persistence(slide, data, index)
        elif data["kind"] == "bridge":
            ppt_bridge(slide, data, index)
        elif data["kind"] == "results":
            ppt_results(slide, data, index)
        else:
            ppt_background(slide)
            ppt_header(slide, data, index)
            ppt_code(slide, data["snippets"])
            ppt_takeaways(slide, data["takeaways"])
        slide.notes_slide.notes_text_frame.text = data["note"]
    presentation.save(PPTX_PATH)


def pdf_rect(page, x, y, width, height, fill, radius=0):
    page.setFillColor(pdf_color(fill))
    if radius:
        page.roundRect(x, y, width, height, radius, fill=1, stroke=0)
    else:
        page.rect(x, y, width, height, fill=1, stroke=0)


def pdf_text(page, text, x, y, width, size, color=TEXT, bold=False, font="Helvetica"):
    font_name = "Helvetica-Bold" if bold and font == "Helvetica" else font
    page.setFont(font_name, size)
    page.setFillColor(pdf_color(color))
    current = y
    for raw_line in text.splitlines():
        line = raw_line
        while len(line) * size * 0.56 > width and " " in line:
            cut = line.rfind(" ", 0, max(1, int(width / (size * 0.56))))
            page.drawString(x, current, line[:cut])
            current -= size * 1.25
            line = line[cut + 1 :]
        page.drawString(x, current, line)
        current -= size * 1.25
    return current


def pdf_header(page, data, index):
    pdf_text(page, data["title"], 40, 505, 760, 23, TEXT, True)
    pdf_text(page, data["subtitle"], 42, 474, 760, 10.5, MUTED)
    pdf_text(page, f"{index:02d}", 902, 504, 30, 9, MUTED, True)


def pdf_takeaways(page, points):
    for index, point in enumerate(points[:2]):
        left = 42 + index * 445
        pdf_rect(page, left, 42, 6, 36, TEAL if index == 0 else ORANGE)
        pdf_text(page, point, left + 15, 69, 410, 10.2, TEXT, True)


def pdf_code_panel(page, snippets, x, y, width, height):
    pdf_rect(page, x, y, width, height, PANEL, 10)
    pdf_text(page, "   |   ".join(item.label for item in snippets), x + 20, y + height - 25, width - 40, 8.2, TEAL, True)
    text = code_text(snippets)
    lines = text.splitlines()
    longest = max(len(line) for line in lines)
    size = max(7.8, min(12.0, (width - 40) / max(longest * 0.60, 1), (height - 58) / max(len(lines) * 1.22, 1)))
    page.setFont("Courier", size)
    page.setFillColor(pdf_color(TEXT))
    current = y + height - 55
    for line in lines:
        page.drawString(x + 20, current, line)
        current -= size * 1.22


def pdf_code(page, snippets):
    line_count = len(code_text(snippets).splitlines())
    height = min(305, max(165, 66 + line_count * 16))
    y = 100 + (335 - height) / 2
    pdf_code_panel(page, snippets, 40, y, 880, height)


def pdf_code_after_flow(page, snippets):
    line_count = len(code_text(snippets).splitlines())
    height = min(248, max(165, 66 + line_count * 13))
    y = 98 + (248 - height) / 2
    pdf_code_panel(page, snippets, 40, y, 880, height)


def pdf_stage_flow(page, labels, palette):
    left = 40
    total_width = 880
    gap = 20
    node_width = (total_width - gap * (len(labels) - 1)) / len(labels)

    for position, (label, color) in enumerate(zip(labels, palette)):
        node_left = left + position * (node_width + gap)
        pdf_rect(page, node_left, 371, node_width, 52, PANEL_LIGHT, 7)
        pdf_rect(page, node_left, 371, 5, 52, color)
        pdf_text(page, label, node_left + 11, 402, node_width - 17, 8.8, TEXT, True)
        if position < len(labels) - 1:
            pdf_text(page, ">", node_left + node_width + 6, 402, 10, 12, ORANGE, True)


def pdf_engine(page, data, index):
    pdf_rect(page, 0, 0, PDF_W, PDF_H, BG)
    pdf_header(page, data, index)
    pdf_stage_flow(
        page,
        ["Command", "execute(command)", "Set / Get / Delete", "state", "buffer"],
        [TEAL, ORANGE, GREEN, TEAL, RED],
    )
    pdf_code_after_flow(page, data["snippets"])
    pdf_takeaways(page, data["takeaways"])


def pdf_persistence(page, data, index):
    pdf_rect(page, 0, 0, PDF_W, PDF_H, BG)
    pdf_header(page, data, index)
    pdf_stage_flow(page, ["state", "buffer Go", "AOF", "snapshot", "restore"], [TEAL, YELLOW, ORANGE, RED, GREEN])
    pdf_code_after_flow(page, data["snippets"])
    pdf_takeaways(page, data["takeaways"])


def pdf_bridge(page, data, index):
    pdf_rect(page, 0, 0, PDF_W, PDF_H, BG)
    pdf_header(page, data, index)
    pdf_stage_flow(
        page,
        ["React", "SDK", "postMessage", "Worker", "wasmRedisExecute", "Engine"],
        [TEAL, YELLOW, ORANGE, ORANGE, GREEN, TEAL],
    )
    pdf_code_after_flow(page, data["snippets"])
    pdf_takeaways(page, data["takeaways"])


def pdf_cover(page, data):
    pdf_rect(page, 0, 0, PDF_W, PDF_H, BG)
    pdf_rect(page, 0, 0, 12, PDF_H, ORANGE)
    pdf_text(page, "WASM / REDIS / REACT", 56, 480, 420, 10, TEAL, True)
    pdf_text(page, data["title"], 55, 412, 520, 45, TEXT, True)
    pdf_text(page, data["subtitle"], 58, 357, 600, 16, MUTED)
    pdf_text(page, "Go en RAM. OPFS sur disque. React reste fluide.", 58, 285, 590, 17, TEXT, True)
    stages = [("React", TEAL), ("SDK TS", YELLOW), ("Worker", ORANGE), ("Go/WASM", GREEN), ("OPFS", RED)]
    for index, (name, color) in enumerate(stages):
        top = 445 - index * 74
        pdf_rect(page, 680, top, 190, 43, color, 8)
        pdf_text(page, name, 700, top + 27, 150, 13, BG, True)
    pdf_text(page, "Projet local, sans serveur", 58, 38, 260, 9, MUTED)


def pdf_flow(page, data, index):
    pdf_rect(page, 0, 0, PDF_W, PDF_H, BG)
    pdf_header(page, data, index)
    stages = [("React", "CRUD + liste", TEAL), ("SDK", "types + Zod", YELLOW), ("Worker", "postMessage", ORANGE), ("Go/WASM", "state + index", GREEN), ("OPFS", "AOF + snapshot", RED)]
    for pos, (name, detail, color) in enumerate(stages):
        left = 42 + pos * 181
        pdf_rect(page, left, 316, 148, 95, PANEL, 8)
        pdf_rect(page, left, 402, 148, 9, color)
        pdf_text(page, name, left + 13, 372, 120, 13, TEXT, True)
        pdf_text(page, detail, left + 13, 342, 120, 8.5, MUTED)
        if pos < 4:
            pdf_text(page, ">", left + 158, 366, 15, 15, ORANGE, True)
    pdf_text(page, "Cycle de persistance", 45, 270, 220, 11, MUTED, True)
    cycle = [("state RAM", TEAL), ("buffer Go", YELLOW), ("AOF ~1 s", ORANGE), ("snapshot ~2 min", RED)]
    for pos, (name, color) in enumerate(cycle):
        left = 45 + pos * 218
        pdf_rect(page, left, 185, 180, 55, PANEL_LIGHT, 7)
        pdf_rect(page, left, 185, 6, 55, color)
        pdf_text(page, name, left + 18, 216, 145, 11, TEXT, True)
        if pos < 3:
            pdf_text(page, ">", left + 194, 216, 15, 14, color, True)
    pdf_takeaways(page, data["takeaways"])


def pdf_results(page, data, index):
    pdf_rect(page, 0, 0, PDF_W, PDF_H, BG)
    pdf_header(page, data, index)
    metrics = [("15,5 x", "gain batch / 30 SET", TEAL), ("~0,3 us", "range selective / 10k", GREEN), ("60 FPS", "scroll mesure", YELLOW), ("29", "lignes DOM / 100k", ORANGE)]
    for pos, (value, label, color) in enumerate(metrics):
        left = 40 + pos * 225
        pdf_rect(page, left, 328, 198, 102, PANEL, 8)
        pdf_text(page, value, left + 16, 390, 165, 20, color, True)
        pdf_text(page, label, left + 16, 350, 165, 8.5, MUTED)
    pdf_text(page, "Demo en 60 secondes", 45, 280, 300, 15, TEXT, True)
    demo = ["1  SET avec TTL, puis Flush", "2  Reload : restore OPFS", "3  Seed 100, filtre >= 50", "4  Editer : 1 -> 2", "5  Lancer Benchmark"]
    for pos, line in enumerate(demo):
        pdf_text(page, line, 50, 245 - pos * 31, 480, 11.5, TEXT, pos == 0)
    pdf_rect(page, 615, 105, 300, 168, PANEL_LIGHT, 8)
    pdf_text(page, "Validation finale", 638, 238, 240, 14, YELLOW, True)
    checks = ["Go race tests", "Tests TypeScript", "Build WASM + React"]
    for pos, check in enumerate(checks):
        pdf_text(page, check, 640, 196 - pos * 35, 190, 11, TEXT, True)
        pdf_text(page, "OK", 853, 196 - pos * 35, 35, 11, GREEN, True)


def build_pdf():
    page = canvas.Canvas(str(PDF_PATH), pagesize=(PDF_W, PDF_H))
    page.setTitle("WasmRedis - presentation")
    for index, data in enumerate(SLIDES, 1):
        if data["kind"] == "cover":
            pdf_cover(page, data)
        elif data["kind"] == "flow":
            pdf_flow(page, data, index)
        elif data["kind"] == "engine":
            pdf_engine(page, data, index)
        elif data["kind"] == "persistence":
            pdf_persistence(page, data, index)
        elif data["kind"] == "bridge":
            pdf_bridge(page, data, index)
        elif data["kind"] == "results":
            pdf_results(page, data, index)
        else:
            pdf_rect(page, 0, 0, PDF_W, PDF_H, BG)
            pdf_header(page, data, index)
            pdf_code(page, data["snippets"])
            pdf_takeaways(page, data["takeaways"])
        page.showPage()
    page.save()


def build_notes():
    lines = ["# Notes orales - WasmRedis", "", "Support de 11 slides pour environ 15 minutes.", ""]
    for index, data in enumerate(SLIDES, 1):
        lines.extend([f"## {index}. {data['title']}", data["note"], ""])
        if data.get("snippets"):
            lines.append("Code affiche :")
            lines.extend(f"- `{item.label}`" for item in data["snippets"])
            lines.append("")
        for point in data.get("takeaways", []):
            lines.append(f"- {point}")
        if data.get("takeaways"):
            lines.append("")
    NOTES_PATH.write_text("\n".join(lines), encoding="utf-8")


def main():
    OUT.mkdir(exist_ok=True)
    build_pptx()
    print(f"PPTX: {PPTX_PATH}")


if __name__ == "__main__":
    main()
